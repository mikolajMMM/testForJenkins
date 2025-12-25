from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


@dataclass(frozen=True)
class Citation:
    n: int
    text: str


CONTENT_PATH = Path("/workspace/paper/thesis.md")


def _add_superscript_citation(paragraph, n: int) -> None:
    run = paragraph.add_run(str(n))
    run.font.superscript = True


def _apply_page_layout(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)


def _apply_base_typography(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


def _format_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1.0)


_CITE_RE = re.compile(r"\{\{cite:([0-9, ]+)\}\}")


def _split_into_blocks(md: str) -> list[str]:
    blocks: list[str] = []
    buf: list[str] = []
    for line in md.splitlines():
        if line.strip() == "":
            if buf:
                blocks.append("\n".join(buf).strip())
                buf = []
            continue
        buf.append(line.rstrip())
    if buf:
        blocks.append("\n".join(buf).strip())
    return blocks


def _parse_citations(text: str) -> tuple[str, list[int]]:
    citations: list[int] = []
    for m in _CITE_RE.finditer(text):
        nums = [int(x.strip()) for x in m.group(1).split(",") if x.strip()]
        citations.extend(nums)
    cleaned = _CITE_RE.sub("", text).strip()
    # de-dup but preserve order
    seen: set[int] = set()
    dedup: list[int] = []
    for n in citations:
        if n not in seen:
            dedup.append(n)
            seen.add(n)
    return cleaned, dedup


def _add_paragraph_with_citations(doc: Document, text: str, citations: list[int]) -> None:
    p = doc.add_paragraph(text)
    _format_paragraph(p)
    if citations:
        p.add_run(" ")
        for i, c in enumerate(citations):
            if i > 0:
                p.add_run(",")
            _add_superscript_citation(p, c)


def _add_heading(doc: Document, title: str, level: int) -> None:
    h = doc.add_heading(title, level=level)
    # headings: no first-line indent; keep 1.5 spacing and left alignment
    h.paragraph_format.first_line_indent = None
    h.paragraph_format.line_spacing = 1.5
    h.paragraph_format.space_after = Pt(6)


def _citations() -> list[Citation]:
    # NOTE: citations are kept intentionally as "primary sources first"
    return [
        Citation(
            1,
            "Malta, Criminal Code (Chapter 9), wersja skonsolidowana (PDF) – Legislation Malta. "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://legislation.mt/eli/cap/9/eng/pdf",
        ),
        Citation(
            2,
            "Malta, Act No. XXII of 2023 (amendment) – Legislation Malta (ELI). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://legislation.mt/eli/act/2023/22/eng",
        ),
        Citation(
            3,
            "Germany, Strafgesetzbuch (StGB) § 218 – dejure.org (źródło pomocnicze; zalecana weryfikacja w źródle urzędowym). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://dejure.org/gesetze/StGB/218.html",
        ),
        Citation(
            4,
            "Germany, Strafgesetzbuch (StGB) § 218a – dejure.org (źródło pomocnicze; zalecana weryfikacja w źródle urzędowym). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://dejure.org/gesetze/StGB/218a.html",
        ),
        Citation(
            5,
            "Germany, Schwangerschaftskonfliktgesetz (SchKG) – dejure.org (źródło pomocnicze; zalecana weryfikacja w źródle urzędowym). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://dejure.org/gesetze/SchKG/",
        ),
        Citation(
            6,
            "Netherlands, Wet afbreking zwangerschap (WAZ), wersja skonsolidowana – wetten.overheid.nl (BWBR0003396). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://wetten.overheid.nl/BWBR0003396/2025-07-05",
        ),
        Citation(
            7,
            "Netherlands, Wet van 22 augustus 2022 tot wijziging van Wet afbreking zwangerschap i.v.m. afschaffen verplichte minimale beraadtermijn (Stb. 2022, 326). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://zoek.officielebekendmakingen.nl/stb-2022-326.html",
        ),
        Citation(
            8,
            "Netherlands, Besluit van 28 september 2022: inwerkingtreding Stb. 2022, 326 per 1 januari 2023 (Stb. 2022, 377). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://zoek.officielebekendmakingen.nl/stb-2022-377.html",
        ),
        Citation(
            9,
            "Netherlands, Wet van 16 januari 2023: legal medicamenteuze afbreking zwangerschap via huisarts (Stb. 2023, 43). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://zoek.officielebekendmakingen.nl/stb-2023-43.html",
        ),
        Citation(
            10,
            "Netherlands, Besluit van 18 mei 2009: wijziging Besluit afbreking zwangerschap (vaststelling duur zwangerschap) (Stb. 2009, 230). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://zoek.officielebekendmakingen.nl/stb-2009-230.html",
        ),
        Citation(
            11,
            "ETPCz, Vo p. Francji, skarga nr 53924/00, wyrok Wielkiej Izby z 08.07.2004 (HUDOC). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://hudoc.echr.coe.int/eng?i=001-61887",
        ),
        Citation(
            12,
            "ETPCz, A, B i C p. Irlandii, skarga nr 25579/05, wyrok Wielkiej Izby z 16.12.2010 (HUDOC). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://hudoc.echr.coe.int/eng?i=001-102332",
        ),
        Citation(
            13,
            "ETPCz, Tysiąc p. Polsce, skarga nr 5410/03, wyrok z 20.03.2007 (HUDOC). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://hudoc.echr.coe.int/eng?i=001-79812",
        ),
        Citation(
            14,
            "ETPCz, R.R. p. Polsce, skarga nr 27617/04, wyrok z 26.05.2011 (HUDOC). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://hudoc.echr.coe.int/eng?i=001-104911",
        ),
        Citation(
            15,
            "ETPCz, P. i S. p. Polsce, skarga nr 57375/08, wyrok z 30.10.2012 (HUDOC). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://hudoc.echr.coe.int/eng?i=001-114098",
        ),
        Citation(
            16,
            "ETPCz, Open Door i Dublin Well Woman p. Irlandii, skarga nr 14234/88, wyrok z 29.10.1992 (HUDOC). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://hudoc.echr.coe.int/eng?i=001-57789",
        ),
        Citation(
            17,
            "BVerfG, Urteil des Zweiten Senats vom 28. Mai 1993, 2 BvF 2/90 u.a. (strona informacyjna BVerfG). "
            f"Dostęp: {date.today().strftime('%d.%m.%Y')}, https://www.bundesverfassungsgericht.de/SharedDocs/Entscheidungen/DE/1993/05/rs19930528_2bvf000290.html",
        ),
    ]


def build_document() -> tuple[Document, list[Citation]]:
    citations = _citations()
    doc = Document()
    _apply_page_layout(doc)
    _apply_base_typography(doc)

    doc.add_paragraph("Temat pracy:")
    doc.add_paragraph(
        "Prawne regulacje przerwania ciąży w wybranych krajach Unii Europejskiej "
        "na przykładzie Malty, Niemiec i Holandii"
    )
    doc.add_paragraph(f"Data opracowania: {date.today().strftime('%d.%m.%Y')}")
    doc.add_page_break()

    md = CONTENT_PATH.read_text(encoding="utf-8")
    for block in _split_into_blocks(md):
        if block.startswith("## "):
            _add_heading(doc, block[3:].strip(), level=1)
            continue
        if block.startswith("### "):
            _add_heading(doc, block[4:].strip(), level=2)
            continue

        text, cites = _parse_citations(block.replace("\n", " ").strip())
        _add_paragraph_with_citations(doc, text, cites)

    doc.add_page_break()
    _add_heading(doc, "Przypisy", level=1)
    for c in citations:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.line_spacing = 1.2
        n_run = p.add_run(f"{c.n}. ")
        n_run.bold = True
        p.add_run(c.text)

    return doc, citations


def main() -> None:
    doc, _ = build_document()
    output_path = "/workspace/paper/Wstep_Rozdzial_I_II_regulacje_aborcji_Malta_Niemcy_Holandia_v2.docx"
    doc.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()

